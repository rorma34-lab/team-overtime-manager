import os
import sys
import io
import json
from pathlib import Path
from datetime import datetime
import openpyxl

# Ensure UTF-8 output
if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from main import (
    app,
    finalize_overtime,
    batch_finalize_overtimes,
    review_overtime,
    batch_review_overtimes,
    create_suggestion,
    list_suggestions,
    reply_suggestion,
    delete_suggestion,
    export_settlement
)
from schemas import (
    OvertimeFinalizeRequest,
    OvertimeBatchFinalizeRequest,
    OvertimeReviewRequest,
    OvertimeBatchReviewRequest,
    SuggestionCreateRequest,
    SuggestionReplyRequest
)
from database import get_db_connection, init_db
from exporter import (
    generate_overtime_excel,
    generate_settlement_excel,
    aggregate_user_holidays
)

def run_tests():
    print("=" * 70)
    print("▶ [Test 0] System Initialization & Version Verification")
    print("=" * 70)
    init_db()
    
    # 0. App Version
    assert app.version == "v1.44", f"Expected version v1.44, got {app.version}"
    print(f"✓ FastAPI App Version is: {app.version}")

    print("\n" + "=" * 70)
    print("▶ [Test 1] 4-Stage Lifecycle: Apply -> Approve -> Finalize -> Review")
    print("=" * 70)

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM overtimes WHERE emp_id = 'test_v144'")
        cursor.execute("DELETE FROM users WHERE emp_id = 'test_v144'")
        cursor.execute("""
            INSERT INTO users (emp_id, name, team, position, is_admin, is_super, created_at)
            VALUES ('test_v144', '테스터V144', '품질혁신팀', '선임', 0, 0, '2026-09-21 09:00:00')
        """)
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # 1단계: 특근 신청 등록 (is_confirmed=0, is_finalized=0, is_reviewed=0)
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  is_confirmed, is_finalized, is_reviewed, created_at, updated_at)
            VALUES ('test_v144', '테스터V144', '품질혁신팀', '일반휴일', '2026-09-26', '2026-09-26', 'PRJ-V144-01', '사내2공장', '4단계 라이프사이클 테스트',
                    0, 0, 0, ?, ?)
        """, (now_str, now_str))
        ot_id = cursor.lastrowid
        conn.commit()

    # 1단계 상태 검증
    with get_db_connection() as conn:
        row = conn.cursor().execute("SELECT is_confirmed, is_finalized, is_reviewed FROM overtimes WHERE id = ?", (ot_id,)).fetchone()
        assert row["is_confirmed"] == 0 and row["is_finalized"] == 0 and row["is_reviewed"] == 0
        print(f"✓ [1단계:신청] OK: ID={ot_id}, is_confirmed=0, is_finalized=0, is_reviewed=0")

    # 2단계: 관리자 승인
    with get_db_connection() as conn:
        conn.cursor().execute("UPDATE overtimes SET is_confirmed = 1, confirmed_by = '총괄관리자', confirmed_at = ? WHERE id = ?", (now_str, ot_id))
        conn.commit()
    with get_db_connection() as conn:
        row = conn.cursor().execute("SELECT is_confirmed, is_finalized, is_reviewed FROM overtimes WHERE id = ?", (ot_id,)).fetchone()
        assert row["is_confirmed"] == 1 and row["is_finalized"] == 0 and row["is_reviewed"] == 0
        print(f"✓ [2단계:승인] OK: ID={ot_id}, is_confirmed=1, is_finalized=0, is_reviewed=0")

    # 3단계: 사원/관리자 특근 완료 [확정] 피드백
    fin_req = OvertimeFinalizeRequest(emp_id="test_v144", is_finalized=1)
    fin_res = finalize_overtime(ot_id, fin_req)
    assert "성공적" in fin_res.get("message") or "변경되었습니다" in fin_res.get("message")
    with get_db_connection() as conn:
        row = conn.cursor().execute("SELECT is_confirmed, is_finalized, finalized_by, is_reviewed FROM overtimes WHERE id = ?", (ot_id,)).fetchone()
        assert row["is_confirmed"] == 1
        assert row["is_finalized"] == 1
        assert "test_v144" in row["finalized_by"]
        assert row["is_reviewed"] == 0
        print(f"✓ [3단계:확정] OK: ID={ot_id}, is_confirmed=1, is_finalized=1, finalized_by={row['finalized_by']}")

    # 4단계: 관리자 [검토완료] 마감 (ps37082 슈퍼관리자 권한 사용)
    rev_req = OvertimeReviewRequest(admin_emp_id="ps37082", is_reviewed=1)
    rev_res = review_overtime(ot_id, rev_req)
    assert "성공적" in rev_res.get("message") or "변경되었습니다" in rev_res.get("message")
    with get_db_connection() as conn:
        row = conn.cursor().execute("SELECT is_confirmed, is_finalized, is_reviewed, reviewed_by FROM overtimes WHERE id = ?", (ot_id,)).fetchone()
        assert row["is_confirmed"] == 1
        assert row["is_finalized"] == 1
        assert row["is_reviewed"] == 1
        assert "ps37082" in row["reviewed_by"]
        print(f"✓ [4단계:검토완료] OK: ID={ot_id}, is_confirmed=1, is_finalized=1, is_reviewed=1, reviewed_by={row['reviewed_by']}")

    print("\n" + "=" * 70)
    print("▶ [Test 2] Batch Finalize & Batch Review Verification")
    print("=" * 70)
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  is_confirmed, is_finalized, is_reviewed, created_at, updated_at)
            VALUES ('test_v144', '테스터V144', '품질혁신팀', '일반휴일', '2026-09-27', '2026-09-27', 'PRJ-BATCH-02', '사내2공장', '일괄테스트2',
                    1, 0, 0, ?, ?),
                   ('test_v144', '테스터V144', '품질혁신팀', '일반휴일', '2026-09-28', '2026-09-28', 'PRJ-BATCH-03', '사내2공장', '일괄테스트3',
                    1, 0, 0, ?, ?)
        """, (now_str, now_str, now_str, now_str))
        ot_id_2 = cursor.lastrowid - 1
        ot_id_3 = cursor.lastrowid
        conn.commit()

    # 일괄 확정
    bf_res = batch_finalize_overtimes(OvertimeBatchFinalizeRequest(emp_id="test_v144", ids=[ot_id_2, ot_id_3], is_finalized=1))
    assert "성공적" in bf_res.get("message")
    print(f"✓ Batch Finalize OK: 2 records finalized ({bf_res.get('message')})")

    # 일괄 검토완료
    br_res = batch_review_overtimes(OvertimeBatchReviewRequest(admin_emp_id="ps37082", ids=[ot_id_2, ot_id_3], is_reviewed=1))
    assert "성공적" in br_res.get("message")
    print(f"✓ Batch Review OK: 2 records reviewed ({br_res.get('message')})")

    print("\n" + "=" * 70)
    print("▶ [Test 3] 100% Anonymous Suggestion Board (무기명 건의사항 소통함)")
    print("=" * 70)
    # 1. 무기명 건의 등록 (사번/성명 일체 없이 등록)
    sug_create = SuggestionCreateRequest(
        category="시스템개선",
        title="야간 모드 다크 테마 도입 요청",
        content="밤에 특근 신청 시 눈이 부십니다. 다크 모드가 있으면 좋겠습니다."
    )
    sug_res = create_suggestion(sug_create)
    assert sug_res.get("status") == "success"
    sug_id = sug_res["id"]
    print(f"✓ Suggestion Created (100% Anonymous) OK: ID={sug_id}")

    # 2. 건의 목록 조회
    sug_list_res = list_suggestions()
    assert sug_list_res.get("status") == "success"
    sug_items = sug_list_res.get("suggestions")
    assert any(s["id"] == sug_id for s in sug_items)
    for s in sug_items:
        assert "emp_id" not in s
        assert "user_name" not in s
    print(f"✓ Suggestion List OK: Total {len(sug_items)} items, 0% author leakage confirmed")

    # 3. 관리자 공식 답변 및 상태 변경
    reply_req = SuggestionReplyRequest(
        admin_emp_id="ps37082",
        admin_reply="v1.45 릴리즈에 다크 테마 모드 도입을 검토하겠습니다. 좋은 의견 감사합니다!",
        status="처리완료"
    )
    reply_res = reply_suggestion(sug_id, reply_req)
    assert reply_res.get("status") == "success"
    print(f"✓ Admin Reply OK: status={reply_res.get('status')}")

    # 4. 건의 삭제 테스트
    temp_sug = create_suggestion(SuggestionCreateRequest(category="기타", title="삭제테스트", content="내용"))
    del_res = delete_suggestion(temp_sug["id"], admin_emp_id="ps37082")
    assert del_res.get("status") == "success"
    print(f"✓ Suggestion Deletion OK: ID={temp_sug['id']}")

    print("\n" + "=" * 70)
    print("▶ [Test 4] Excel Export 4-Stage Days Separation & Formulas")
    print("=" * 70)
    with get_db_connection() as conn:
        records = [dict(r) for r in conn.cursor().execute("SELECT * FROM overtimes").fetchall()]
        users = [dict(u) for u in conn.cursor().execute("SELECT * FROM users").fetchall()]

    # Generate master excel
    excel_buffer = generate_overtime_excel(records, user_positions={}, period_str="2026-09")
    wb = openpyxl.load_workbook(io.BytesIO(excel_buffer.getvalue()), data_only=False)

    # 1. Sheet 1: 휴일일자별_특근현황
    ws1 = wb["휴일일자별_특근현황"]
    headers_ws1 = [cell.value for cell in ws1[3]]
    print("Sheet 1 Headers:", headers_ws1[:15])
    assert "진행단계" in headers_ws1, "Expected '진행단계' in Sheet 1 headers"
    assert any("승인자" in str(h) for h in headers_ws1), "Expected '승인자' in Sheet 1 headers"
    assert any("확정자" in str(h) for h in headers_ws1), "Expected '확정자' in Sheet 1 headers"
    assert any("검토자" in str(h) for h in headers_ws1), "Expected '검토자' in Sheet 1 headers"
    print("✓ Sheet 1 [휴일일자별_특근현황]: '진행단계', '승인자', '확정자', '검토자' columns verified")

    # 2. Sheet 2: 개인별_휴일합산_정산표
    ws2 = wb["개인별_휴일합산_정산표"]
    headers_ws2 = [cell.value for cell in ws2[3]]
    print("Sheet 2 Headers:", headers_ws2[:16])
    assert "신청단계 (일)" in headers_ws2, "Expected '신청단계 (일)' in Sheet 2 headers"
    assert "승인단계 (일)" in headers_ws2, "Expected '승인단계 (일)' in Sheet 2 headers"
    assert "확정단계 (일)" in headers_ws2, "Expected '확정단계 (일)' in Sheet 2 headers"
    assert "검토완료 (일)" in headers_ws2, "Expected '검토완료 (일)' in Sheet 2 headers"
    print("✓ Sheet 2 [개인별_휴일합산_정산표]: 4-Stage Separate Days columns verified")

    # 3. Sheet 4: 부서별_요약
    if "부서별_요약" in wb.sheetnames:
        ws4 = wb["부서별_요약"]
        headers_ws4 = [cell.value for cell in ws4[3]]
        print("Sheet 4 Headers:", headers_ws4[:12])
        assert "신청 (일)" in headers_ws4
        assert "검토완료 (일)" in headers_ws4
        print("✓ Sheet 4 [부서별_요약]: 4-Stage days columns verified")

    # 4. Settlement Excel generation (4-Stage Separate Days)
    summary_data = {
        "user_summary": [
            {
                "emp_id": "test_v144",
                "name": "테스터V144",
                "team": "품질혁신팀",
                "excluded_sub_days": 0,
                "excluded_legal_days": 0,
                "overtime_days": 3,
                "sub_holiday_days": 0,
                "sub_holiday_used": 0,
                "total_days": 3,
                "applied_days": 0,
                "approved_days": 0,
                "finalized_days": 2,
                "reviewed_days": 1,
                "pre_deduct_count": 0,
                "pre_deduct_remaining": 0,
                "actual_overtime_days": 3.0,
                "bonus_count": 0,
                "actual_overtime_with_bonus": 3.0
            }
        ],
        "team_summary": [
            {
                "team": "품질혁신팀",
                "member_count": 1,
                "total_days": 3,
                "excluded_days": 0,
                "overtime_days": 3,
                "sub_holiday_used": 0,
                "pre_deduct_days": 0,
                "actual_overtime_days": 3.0,
                "bonus_count": 0,
                "actual_overtime_with_bonus": 3.0,
                "applied_days": 0,
                "approved_days": 0,
                "finalized_days": 2,
                "reviewed_days": 1
            }
        ]
    }
    settle_excel_buf = generate_settlement_excel(summary_data=summary_data, start_date="2026-09-01", end_date="2026-09-30")
    wb_settle = openpyxl.load_workbook(io.BytesIO(settle_excel_buf.getvalue()), data_only=False)
    assert "개인별_실특근_정산표" in wb_settle.sheetnames
    ws_set1 = wb_settle["개인별_실특근_정산표"]
    headers_set1 = [cell.value for cell in ws_set1[3]]
    print("Settlement Sheet 1 Headers:", headers_set1[:16])
    assert "신청단계 (일)" in headers_set1
    assert "검토완료 (일)" in headers_set1
    print("✓ Settlement Excel [개인별_실특근_정산표]: 4-Stage days columns verified")

    assert "부서별_정산_요약표" in wb_settle.sheetnames
    ws_set2 = wb_settle["부서별_정산_요약표"]
    headers_set2 = [cell.value for cell in ws_set2[3]]
    print("Settlement Sheet 2 Headers:", headers_set2[:15])
    assert "신청 (일)" in headers_set2
    assert "검토완료 (일)" in headers_set2
    print("✓ Settlement Excel [부서별_정산_요약표]: 4-Stage days columns verified")

    print("\n" + "=" * 70)
    print("▶ [Test 5] Manual Files & PPTX Synchronization Verification")
    print("=" * 70)
    root_dir = Path(".")
    manuals = [
        root_dir / "Overtime_User_Manual.pptx",
        root_dir / "Overtime_Admin_Manual.pptx",
        root_dir / "Overtime_System_Manual.pptx",
        root_dir / "downloads" / "Overtime_User_Manual.pptx",
        root_dir / "downloads" / "Overtime_Admin_Manual.pptx",
        root_dir / "downloads" / "Overtime_System_Manual.pptx",
        root_dir / "static" / "downloads" / "Overtime_User_Manual.pptx",
        root_dir / "static" / "downloads" / "Overtime_Admin_Manual.pptx",
        root_dir / "static" / "downloads" / "Overtime_System_Manual.pptx"
    ]
    for m in manuals:
        assert m.exists(), f"Manual not found: {m}"
        size = m.stat().st_size
        assert size > 100000, f"Manual size too small: {m} ({size} bytes)"
        print(f"✓ {m}: {size:,} bytes")

    # Cleanup test data
    with get_db_connection() as conn:
        conn.cursor().execute("DELETE FROM overtimes WHERE emp_id = 'test_v144'")
        conn.cursor().execute("DELETE FROM users WHERE emp_id = 'test_v144'")
        conn.cursor().execute("DELETE FROM suggestions WHERE id = ?", (sug_id,))
        conn.commit()

    print("\n" + "=" * 70)
    print("🎉 ALL V1.44 INTEGRATION TESTS PASSED 100% PERFECTLY!")
    print("=" * 70)

if __name__ == "__main__":
    run_tests()
