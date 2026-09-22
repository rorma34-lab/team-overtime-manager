import io
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ===========================================================================
#  공통 스타일 팩토리 (openpyxl 엔진 전용)
# ===========================================================================

FONT_FAMILY = "맑은 고딕"

def _make_styles():
    """공통 스타일 dict 반환 (4단계 특근 라이프사이클 스타일 포함)"""
    thin = Side(style='thin', color='CBD5E1')
    medium = Side(style='medium', color='94A3B8')

    return {
        "title_font":    Font(name=FONT_FAMILY, size=16, bold=True, color="0F172A"),
        "subtitle_font": Font(name=FONT_FAMILY, size=9,  italic=True, color="64748B"),
        "header_font":   Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF"),
        "data_font":     Font(name=FONT_FAMILY, size=9,  color="1E293B"),
        "total_font":    Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A"),
        "hl_font":       Font(name=FONT_FAMILY, size=9,  bold=True, color="0D9488"),  # 실특근 강조
        "conf_font":     Font(name=FONT_FAMILY, size=9,  bold=True, color="15803D"),
        "pend_font":     Font(name=FONT_FAMILY, size=9,  color="B45309"),

        # 4단계 상태별 폰트 & 배경색
        "font_stage_apply": Font(name=FONT_FAMILY, size=9, bold=True, color="B45309"),  # 신청: 주황
        "font_stage_appr":  Font(name=FONT_FAMILY, size=9, bold=True, color="1D4ED8"),  # 승인: 파랑
        "font_stage_fin":   Font(name=FONT_FAMILY, size=9, bold=True, color="047857"),  # 확정: 에메랄드초록
        "font_stage_rev":   Font(name=FONT_FAMILY, size=9, bold=True, color="6D28D9"),  # 검토완료: 보라

        "fill_stage_apply": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),  # 연노랑
        "fill_stage_appr":  PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),  # 연파랑
        "fill_stage_fin":   PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),  # 연민트/초록
        "fill_stage_rev":   PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),  # 연보라

        "fill_navy":    PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"),
        "fill_teal":    PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid"),
        "fill_slate":   PatternFill(start_color="334155", end_color="334155", fill_type="solid"),
        "fill_purple":  PatternFill(start_color="5B21B6", end_color="5B21B6", fill_type="solid"),
        "fill_alt":     PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        "fill_total":   PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"),
        "fill_conf":    PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "fill_pend":    PatternFill(start_color="D97706", end_color="D97706", fill_type="solid"),  # 신청단계 헤더: 호박/주황(흰색 글씨와 고대비 보장)
        "fill_header_apply": PatternFill(start_color="D97706", end_color="D97706", fill_type="solid"),
        "fill_hl":      PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid"),
        "fill_gold":    PatternFill(start_color="D97706", end_color="D97706", fill_type="solid"),
        "fill_none":    PatternFill(fill_type=None),

        "border_thin":   Border(left=thin, right=thin, top=thin, bottom=thin),
        "border_total":  Border(left=medium, right=medium, top=medium,
                                bottom=Side(style='double', color='0F172A')),
        "border_header": Border(left=medium, right=medium, top=medium, bottom=medium),

        "align_center": Alignment(horizontal="center", vertical="center"),
        "align_left":   Alignment(horizontal="left",   vertical="center", indent=1),
        "align_right":  Alignment(horizontal="right",  vertical="center"),
        "align_wrap":   Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _auto_col_width(ws, skip_rows=(1, 2)):
    """모든 시트 컬럼 너비 자동 조정 (타이틀 행 제외)"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in skip_rows:
                continue
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                val_str = "999999"
            val_len = sum(2 if ord(c) > 128 else 1 for c in val_str)
            if val_len > max_len:
                max_len = val_len
        ws.column_dimensions[col_letter].width = max(max_len + 4, 10)


def _write_title(ws, text: str, subtitle: str, merge_cols: int, s: dict):
    """타이틀 + 서브타이틀 2행 공통 기입"""
    ws.merge_cells(f"A1:{get_column_letter(merge_cols)}1")
    tc = ws["A1"]
    tc.value = text
    tc.font = s["title_font"]
    tc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{get_column_letter(merge_cols)}2")
    sc = ws["A2"]
    sc.value = subtitle
    sc.font = s["subtitle_font"]
    sc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18


def _write_header_row(ws, row: int, headers: list, fills: list, s: dict):
    """헤더 행 기입"""
    ws.row_dimensions[row].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = s["header_font"]
        cell.fill = fills[col_idx - 1] if fills[col_idx - 1] else s["fill_navy"]
        cell.alignment = s["align_wrap"]
        cell.border = s["border_header"]


def _write_data_row(ws, row: int, values: list, s: dict, is_even: bool,
                    center_cols: set, right_cols: set,
                    num_fmt_map: dict = None, hl_cols: set = None):
    """데이터 행 기입"""
    fill = s["fill_alt"] if is_even else s["fill_none"]
    ws.row_dimensions[row].height = 22
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = s["data_font"]
        cell.fill = fill
        cell.border = s["border_thin"]
        if col_idx in center_cols:
            cell.alignment = s["align_center"]
        elif col_idx in right_cols:
            cell.alignment = s["align_right"]
        else:
            cell.alignment = s["align_left"]
        if num_fmt_map and col_idx in num_fmt_map:
            cell.number_format = num_fmt_map[col_idx]
        if hl_cols and col_idx in hl_cols:
            cell.font = s["hl_font"]


def _write_total_row(ws, row: int, label: str, label_merge_end: int,
                     sum_cols: list, s: dict, data_start_row: int = 4):
    """합계 행 기입 (SUM 수식 + 더블 밑줄 테두리)"""
    ws.row_dimensions[row].height = 26
    if label_merge_end > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=label_merge_end)
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = s["total_font"]
    lc.alignment = s["align_center"]
    for c in range(1, label_merge_end + 1):
        ws.cell(row=row, column=c).fill = s["fill_total"]
        ws.cell(row=row, column=c).border = s["border_total"]

    for col_i, num_fmt, is_hl in sum_cols:
        f_val = f"=SUM({get_column_letter(col_i)}{data_start_row}:{get_column_letter(col_i)}{row - 1})"
        c = ws.cell(row=row, column=col_i, value=f_val)
        c.font = s["hl_font"] if is_hl else s["total_font"]
        c.fill = s["fill_hl"] if is_hl else s["fill_total"]
        c.border = s["border_total"]
        c.alignment = s["align_right"]
        c.number_format = num_fmt


def _apply_stage_cell_style(cell, stage_text: str, s: dict):
    """진행단계별 셀 배경/글자 스타일 적용"""
    if stage_text == "검토완료":
        cell.fill = s["fill_stage_rev"]
        cell.font = s["font_stage_rev"]
    elif stage_text == "확정":
        cell.fill = s["fill_stage_fin"]
        cell.font = s["font_stage_fin"]
    elif stage_text == "승인":
        cell.fill = s["fill_stage_appr"]
        cell.font = s["font_stage_appr"]
    else:  # 신청
        cell.fill = s["fill_stage_apply"]
        cell.font = s["font_stage_apply"]


# ===========================================================================
#  보조 데이터 처리 함수
# ===========================================================================

def expand_records_by_holiday_date(records: list) -> list:
    """다일 특근 신청 내역을 개별 휴일 날짜 단위로 전개 및 4단계 진행정보 부여"""
    WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]
    expanded = []

    for r in records:
        s_date_str = r.get("start_date", "")
        e_date_str = r.get("end_date", "") or s_date_str

        try:
            d1 = datetime.strptime(s_date_str, "%Y-%m-%d")
            d2 = datetime.strptime(e_date_str, "%Y-%m-%d")
            if d2 < d1:
                d2 = d1
        except Exception:
            d1 = datetime.now()
            d2 = d1

        total_days = (d2 - d1).days + 1
        sub_used = float(r.get("sub_holiday_used", 0.0) or 0.0)
        sub_date = r.get("sub_holiday_date", "") or "-"

        is_conf = bool(r.get("is_confirmed", 0))
        is_fin = bool(r.get("is_finalized", 0))
        is_rev = bool(r.get("is_reviewed", 0))

        if is_rev:
            stage_text = "검토완료"
            stage_order = 4
        elif is_fin:
            stage_text = "확정"
            stage_order = 3
        elif is_conf:
            stage_text = "승인"
            stage_order = 2
        else:
            stage_text = "신청"
            stage_order = 1

        curr = d1
        day_idx = 0
        while curr <= d2:
            date_str = curr.strftime("%Y-%m-%d")
            w_idx = curr.weekday()
            w_str = WEEKDAYS[w_idx]

            cur_sub_date = sub_date if (day_idx == 0 and sub_date != "-") else ("-" if total_days > 1 else sub_date)
            cur_sub_used = sub_used if (day_idx == 0) else 0.0
            cur_is_pre = int(r.get("is_pre_deduct", 0) or 0) if day_idx == 0 else 0

            expanded.append({
                "holiday_date": date_str,
                "weekday": w_str,
                "emp_id": r.get("emp_id", ""),
                "user_name": r.get("user_name", ""),
                "team": r.get("team", ""),
                "category": r.get("category", ""),
                "bonus_granted": r.get("bonus_granted", 0),
                "bonus_text": "부여(O)" if r.get("bonus_granted", 0) == 1 else "미부여(-)",
                "stage_text": stage_text,
                "stage_order": stage_order,
                "holiday_days": 1,
                "sub_holiday_date": cur_sub_date,
                "sub_holiday_used": cur_sub_used,
                "is_pre_deduct": cur_is_pre,
                "pre_deduct_text": "O" if cur_is_pre else "-",
                "trip_start_date": r.get("trip_start_date", "") or "-" if day_idx == 0 else "-",
                "trip_end_date": r.get("trip_end_date", "") or "-" if day_idx == 0 else "-",
                "project_no": r.get("project_no", "") or "-",
                "location": r.get("location", "") or "-",
                "reason": r.get("reason", "") or "-",
                "is_confirmed": is_conf,
                "confirmed_by": r.get("confirmed_by", "") or "-",
                "is_finalized": is_fin,
                "finalized_by": r.get("finalized_by", "") or "-",
                "finalized_at": r.get("finalized_at", "") or "-",
                "is_reviewed": is_rev,
                "reviewed_by": r.get("reviewed_by", "") or "-",
                "reviewed_at": r.get("reviewed_at", "") or "-",
                "created_at": r.get("created_at", "")
            })
            curr += timedelta(days=1)
            day_idx += 1

    expanded.sort(key=lambda x: (x["holiday_date"], x["team"], x["user_name"]))
    return expanded


def aggregate_user_holidays(records: list, user_positions: dict = None) -> list:
    """개인별 휴일수 및 단계별(신청,승인,확정,검토완료) 실특근일+보너스 세부 합산 집계"""
    user_positions = user_positions or {}
    user_map = {}
    user_trips = {}  # emp_id -> [(start, end)]

    # 1차: 대체휴무 출장기간 수집
    for r in records:
        emp_id = r.get("emp_id", "")
        cat = (r.get("category", "") or "").strip()
        t_s = (r.get("trip_start_date", "") or "").strip()
        t_e = (r.get("trip_end_date", "") or "").strip()
        if cat in ["대체휴무", "대체휴일"] and t_s and t_e:
            if emp_id not in user_trips:
                user_trips[emp_id] = []
            user_trips[emp_id].append((t_s, t_e))

    for r in records:
        emp_id = r.get("emp_id", "")
        name = r.get("user_name", "")
        team = r.get("team", "")
        cat = r.get("category", "")
        if cat == "대체휴일":
            cat = "대체휴무"
        sub_used = float(r.get("sub_holiday_used", 0.0) or 0.0)
        is_pre = int(r.get("is_pre_deduct", 0) or 0)
        s_date = r.get("start_date", "")
        e_date = r.get("end_date") or s_date

        try:
            d1 = datetime.strptime(s_date, "%Y-%m-%d")
            d2 = datetime.strptime(e_date, "%Y-%m-%d")
            days = max(1, (d2 - d1).days + 1)
        except Exception:
            days = 1

        is_conf = bool(r.get("is_confirmed", 0))
        is_fin = bool(r.get("is_finalized", 0))
        is_rev = bool(r.get("is_reviewed", 0))
        is_bonus = 1 if int(r.get("bonus_granted", 0) or 0) == 1 else 0

        if emp_id not in user_map:
            user_map[emp_id] = {
                "emp_id": emp_id,
                "name": name,
                "team": team,
                "position": user_positions.get(emp_id, "-"),
                "sub_work_days": 0,
                "legal_holiday_days": 0,
                "normal_holiday_days": 0,
                "sub_holiday_days": 0,
                "pre_deduct_count": 0,
                "trip_pre_deduct_count": 0,
                "total_days": 0,
                "sub_holiday_used": 0.0,
                "actual_overtime_days": 0.0,
                "bonus_count": 0,
                "records_count": 0,
                # 4단계별 총 특근일수 (대체/법정/일반 포함)
                "applied_days": 0,
                "approved_days": 0,
                "finalized_days": 0,
                "reviewed_days": 0,
                # 4단계별 일반특근일수 (실특근 대상)
                "applied_ot_days": 0,
                "approved_ot_days": 0,
                "finalized_ot_days": 0,
                "reviewed_ot_days": 0,
                # 4단계별 보너스 건수
                "applied_bonus": 0,
                "approved_bonus": 0,
                "finalized_bonus": 0,
                "reviewed_bonus": 0,
                # 4단계별 최종 실특근일+보너스 [일]
                "applied_final_with_bonus": 0.0,
                "approved_final_with_bonus": 0.0,
                "finalized_final_with_bonus": 0.0,
                "reviewed_final_with_bonus": 0.0,
                "actual_overtime_with_bonus": 0.0
            }

        u = user_map[emp_id]
        u["records_count"] += 1
        u["sub_holiday_used"] += sub_used

        if cat == "대체근무":
            u["sub_work_days"] += days
        elif cat == "법정휴일":
            u["legal_holiday_days"] += days
        elif cat == "대체휴무":
            u["sub_holiday_days"] += days
        else:  # 일반휴일 + 기타
            u["normal_holiday_days"] += days

        # 4단계 일수 및 일반특근/보너스 누적 (검토완료 > 확정 > 승인 > 신청)
        if is_rev:
            u["reviewed_days"] += days
            if cat == "일반휴일":
                u["reviewed_ot_days"] += days
            if is_bonus:
                u["reviewed_bonus"] += 1
        elif is_fin:
            u["finalized_days"] += days
            if cat == "일반휴일":
                u["finalized_ot_days"] += days
            if is_bonus:
                u["finalized_bonus"] += 1
        elif is_conf:
            u["approved_days"] += days
            if cat == "일반휴일":
                u["approved_ot_days"] += days
            if is_bonus:
                u["approved_bonus"] += 1
        else:
            u["applied_days"] += days
            if cat == "일반휴일":
                u["applied_ot_days"] += days
            if is_bonus:
                u["applied_bonus"] += 1

        # 총 특근일수 = 대체근무 + 법정휴일 + 일반휴일 (대체휴무 제외)
        u["total_days"] = int(u["sub_work_days"] + u["legal_holiday_days"] + u["normal_holiday_days"])

        # 총 보너스 개수 집계
        if is_bonus:
            u["bonus_count"] += 1

        if is_pre:
            u["pre_deduct_count"] += 1
            trips = user_trips.get(emp_id, [])
            for ts, te in trips:
                if not (e_date < ts or s_date > te):
                    u["trip_pre_deduct_count"] += 1
                    break

    # 2차: 개인별 차감 분배 및 단계별 최종 실특근일+보너스 확정 계산
    for u in user_map.values():
        u["pre_deduct_remaining"] = max(0, u["pre_deduct_count"] - u["trip_pre_deduct_count"])
        net_deduct = float(u["pre_deduct_count"] + max(0.0, float(u["sub_holiday_days"] - u["trip_pre_deduct_count"])))
        u["actual_overtime_days"] = max(0.0, round(float(u["normal_holiday_days"] - net_deduct), 1))
        u["actual_overtime_with_bonus"] = max(0.0, round(float(u["actual_overtime_days"] + u["bonus_count"]), 1))

        # 차감의 단계별 배분 (우선순위: 검토완료 -> 확정 -> 승인 -> 신청 순)
        rem_ded = max(0.0, net_deduct)

        d_rev = min(float(u["reviewed_ot_days"]), rem_ded)
        rev_net = max(0.0, float(u["reviewed_ot_days"]) - d_rev)
        rem_ded = max(0.0, rem_ded - d_rev)

        d_fin = min(float(u["finalized_ot_days"]), rem_ded)
        fin_net = max(0.0, float(u["finalized_ot_days"]) - d_fin)
        rem_ded = max(0.0, rem_ded - d_fin)

        d_appr = min(float(u["approved_ot_days"]), rem_ded)
        appr_net = max(0.0, float(u["approved_ot_days"]) - d_appr)
        rem_ded = max(0.0, rem_ded - d_appr)

        d_apply = min(float(u["applied_ot_days"]), rem_ded)
        apply_net = max(0.0, float(u["applied_ot_days"]) - d_apply)
        rem_ded = max(0.0, rem_ded - d_apply)

        # 각 단계별 최종 실특근일+보너스 [일]
        u["applied_final_with_bonus"] = max(0.0, round(apply_net + u["applied_bonus"], 1))
        u["approved_final_with_bonus"] = max(0.0, round(appr_net + u["approved_bonus"], 1))
        u["finalized_final_with_bonus"] = max(0.0, round(fin_net + u["finalized_bonus"], 1))
        u["reviewed_final_with_bonus"] = max(0.0, round(rev_net + u["reviewed_bonus"], 1))

        # 검증: 단계별 합산과 전체 합계 동기화
        u["actual_overtime_with_bonus"] = max(0.0, round(
            u["applied_final_with_bonus"] + u["approved_final_with_bonus"] +
            u["finalized_final_with_bonus"] + u["reviewed_final_with_bonus"], 1
        ))

    user_list = list(user_map.values())
    user_list.sort(key=lambda x: (x["team"], x["name"]))
    return user_list


def aggregate_team_holidays(user_summaries: list) -> list:
    """부서(팀)별 합산 요약 집계 (4단계 및 단계별 최종 실특근일+보너스 포함)"""
    team_map = {}
    for u in user_summaries:
        t = u["team"]
        if t not in team_map:
            team_map[t] = {
                "team": t,
                "member_count": 0,
                "sub_work_days": 0,
                "legal_holiday_days": 0,
                "normal_holiday_days": 0,
                "sub_holiday_days": 0,
                "total_days": 0,
                "actual_overtime_days": 0.0,
                "bonus_count": 0,
                "applied_days": 0,
                "approved_days": 0,
                "finalized_days": 0,
                "reviewed_days": 0,
                "applied_final_with_bonus": 0.0,
                "approved_final_with_bonus": 0.0,
                "finalized_final_with_bonus": 0.0,
                "reviewed_final_with_bonus": 0.0,
                "actual_overtime_with_bonus": 0.0,
                "records_count": 0
            }
        tm = team_map[t]
        tm["member_count"] += 1
        tm["sub_work_days"] += u["sub_work_days"]
        tm["legal_holiday_days"] += u["legal_holiday_days"]
        tm["normal_holiday_days"] += u["normal_holiday_days"]
        tm["sub_holiday_days"] += u.get("sub_holiday_days", 0)
        tm["total_days"] += u["total_days"]
        tm["actual_overtime_days"] = max(0.0, round(tm["actual_overtime_days"] + u["actual_overtime_days"], 1))
        tm["bonus_count"] += u.get("bonus_count", 0)
        tm["records_count"] += u["records_count"]
        tm["applied_days"] += u.get("applied_days", 0)
        tm["approved_days"] += u.get("approved_days", 0)
        tm["finalized_days"] += u.get("finalized_days", 0)
        tm["reviewed_days"] += u.get("reviewed_days", 0)
        tm["applied_final_with_bonus"] = max(0.0, round(tm["applied_final_with_bonus"] + u.get("applied_final_with_bonus", 0.0), 1))
        tm["approved_final_with_bonus"] = max(0.0, round(tm["approved_final_with_bonus"] + u.get("approved_final_with_bonus", 0.0), 1))
        tm["finalized_final_with_bonus"] = max(0.0, round(tm["finalized_final_with_bonus"] + u.get("finalized_final_with_bonus", 0.0), 1))
        tm["reviewed_final_with_bonus"] = max(0.0, round(tm["reviewed_final_with_bonus"] + u.get("reviewed_final_with_bonus", 0.0), 1))
        tm["actual_overtime_with_bonus"] = max(0.0, round(tm["actual_overtime_with_bonus"] + u.get("actual_overtime_with_bonus", 0.0), 1))

    team_list = list(team_map.values())
    team_list.sort(key=lambda x: x["team"])
    return team_list


# ===========================================================================
#  메인 엑셀 생성 함수 1: 특근 신청 내역 전체 (4개 시트)
# ===========================================================================

def generate_overtime_excel(records: list, user_positions: dict = None, period_str: str = "") -> io.BytesIO:
    """
    특근 목록 데이터를 openpyxl로 4개 시트 워크북 생성 (4단계 라이프사이클 구분 지원)
    - 시트1: [휴일일자별_특근현황]  개별 날짜 전개 + 진행단계 + 승인/확정/검토자
    - 시트2: [개인별_휴일합산_정산표]  인원별 분류별 집계 + 신청/승인/확정/검토 4단계 구분 집계 + SUM 공식
    - 시트3: [특근신청_전체원장]  원장 (4단계 진행상태 표시)
    - 시트4: [부서별_요약]  팀별 집계 + 4단계 일수
    """
    s = _make_styles()
    wb = Workbook()
    now_str = datetime.now().strftime("%Y년 %m월 %d일 %H:%M")

    if not period_str and records:
        all_starts = [r.get("start_date") for r in records if r.get("start_date")]
        all_ends = [r.get("end_date") or r.get("start_date") for r in records if r.get("start_date")]
        if all_starts and all_ends:
            min_s = min(all_starts)
            max_e = max(all_ends)
            period_str = f"{min_s} ~ {max_e}"
    if not period_str:
        period_str = "전체 기간"

    # ─────────────────────────────────────────────
    # 시트 1: 휴일일자별_특근현황
    # ─────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "휴일일자별_특근현황"

    _write_title(ws1,
                 "일자별 휴일 특근(초과근무) 상세 현황 (4단계 진행상태)",
                 f"취합 기간: {period_str}  |  출력일시: {now_str}  |  단계: 신청 → 승인 → 확정 → 검토완료  |  일수 단위: 순수 숫자",
                 22, s)

    headers1 = [
        "순번", "휴일날짜", "요일", "사번", "성명", "소속팀", "특근분류", "보너스 부여",
        "진행단계", "휴일일수", "대체휴무 사용일", "대체휴가 사용일수", "사전차감",
        "출장기간(시작)", "출장기간(종료)",
        "프로젝트 번호", "근무 장소", "특근 사유", "승인(승인자)", "확정(확정자)", "검토(검토자)", "신청일시"
    ]
    fills1 = [None] * len(headers1)
    fills1[8] = s["fill_slate"]   # 진행단계
    _write_header_row(ws1, 3, headers1, fills1, s)
    ws1.freeze_panes = "A4"

    expanded = expand_records_by_holiday_date(records)
    r_idx = 4
    total_hdays = 0
    total_sub_used_sum = 0.0

    for idx, r in enumerate(expanded, 1):
        h_val = int(r["holiday_days"])
        sub_v = float(r["sub_holiday_used"])
        total_hdays += h_val
        total_sub_used_sum += sub_v

        fin_info = f"{r['finalized_by']}" if r["is_finalized"] and r["finalized_by"] != "-" else ("완료" if r["is_finalized"] else "-")
        rev_info = f"{r['reviewed_by']}" if r["is_reviewed"] and r["reviewed_by"] != "-" else ("완료" if r["is_reviewed"] else "-")
        conf_info = r["confirmed_by"] if r["is_confirmed"] and r["confirmed_by"] != "-" else ("승인대기" if not r["is_confirmed"] else "승인완료")

        row_data = [
            idx,
            r["holiday_date"],
            r["weekday"],
            r["emp_id"],
            r["user_name"],
            r["team"],
            r["category"],
            r.get("bonus_text", "미부여(-)"),
            r["stage_text"],
            h_val,
            r["sub_holiday_date"],
            sub_v,
            r.get("pre_deduct_text", "-"),
            r.get("trip_start_date", "-"),
            r.get("trip_end_date", "-"),
            r["project_no"],
            r["location"],
            r["reason"],
            conf_info,
            fin_info,
            rev_info,
            r["created_at"]
        ]

        center_c = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 13, 14, 15, 19, 20, 21, 22}
        right_c  = {10, 12}
        num_fmt  = {10: "#,##0", 12: "0.0"}

        _write_data_row(ws1, r_idx, row_data, s, r_idx % 2 == 0, center_c, right_c, num_fmt)

        # 진행단계 셀 스타일 적용
        stage_cell = ws1.cell(row=r_idx, column=9)
        _apply_stage_cell_style(stage_cell, r["stage_text"], s)

        r_idx += 1

    # 합계 행
    _write_total_row(ws1, r_idx, f"총 합계 ({len(expanded)}건)", 9,
                     [(10, "#,##0", False), (12, "0.0", False)], s)
    ws1.cell(row=r_idx, column=11, value="-").alignment = s["align_center"]
    ws1.cell(row=r_idx, column=11).fill = s["fill_total"]
    ws1.cell(row=r_idx, column=11).border = s["border_total"]
    for c in range(13, 23):
        ws1.cell(row=r_idx, column=c, value="").fill = s["fill_total"]
        ws1.cell(row=r_idx, column=c).border = s["border_total"]

    # ─────────────────────────────────────────────
    # 시트 2: 개인별_휴일합산_정산표 (4단계 구분 명시)
    # ─────────────────────────────────────────────
    ws2 = wb.create_sheet(title="개인별_휴일합산_정산표")
    _write_title(ws2,
                 "개인별 특근 휴일수 세부 합산 및 단계별(신청/승인/확정/검토) 실특근일 정산표",
                 f"취합 기간: {period_str}  |  취합 일시: {now_str}  |  [총특근일수] = [대체근무]+[법정휴일]+[일반휴일]  |  [★최종 실특근일] = [일반특근] - [사전차감] - ([대체휴무] - [출장내사전차감])",
                 23, s)

    headers2 = [
        "순번", "사원번호", "성명", "소속팀", "직급",
        "대체근무 (일)", "법정휴일 (일)", "일반휴일 (일)", "대체휴무 (일)",
        "총 특근일수\n(대체+법정+일반)",
        "신청단계 (일)", "승인단계 (일)", "확정단계 (일)", "검토완료 (일)",
        "총 사전차감 (회)", "사전차감 잔여수\n(총사전차감 - 출장내사전차감)",
        "보너스 부여 (건)",
        "신청: 실특근+보너스 [일]", "승인: 실특근+보너스 [일]",
        "확정: 실특근+보너스 [일]", "검토완료: 실특근+보너스 [일]",
        "★ 최종 실특근일+보너스\n(전체합계) [일]", "신청건수"
    ]
    fills2 = [
        None, None, None, None, None,
        s["fill_slate"], s["fill_slate"], s["fill_navy"], s["fill_navy"],
        s["fill_navy"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_teal"], s["fill_teal"],
        s["fill_gold"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_navy"], None
    ]
    _write_header_row(ws2, 3, headers2, fills2, s)
    ws2.freeze_panes = "A4"

    user_summaries = aggregate_user_holidays(records, user_positions)
    r2 = 4
    for idx, u in enumerate(user_summaries, 1):
        row_data = [
            idx,
            u["emp_id"],
            u["name"],
            u["team"],
            u["position"],
            int(u["sub_work_days"]),
            int(u["legal_holiday_days"]),
            int(u["normal_holiday_days"]),
            int(u.get("sub_holiday_days", 0)),
            int(u["total_days"]),
            int(u.get("applied_days", 0)),
            int(u.get("approved_days", 0)),
            int(u.get("finalized_days", 0)),
            int(u.get("reviewed_days", 0)),
            int(u.get("pre_deduct_count", 0)),
            int(u.get("pre_deduct_remaining", 0)),
            int(u.get("bonus_count", 0)),
            float(u.get("applied_final_with_bonus", 0.0)),
            float(u.get("approved_final_with_bonus", 0.0)),
            float(u.get("finalized_final_with_bonus", 0.0)),
            float(u.get("reviewed_final_with_bonus", 0.0)),
            float(u.get("actual_overtime_with_bonus", 0.0)),
            int(u["records_count"])
        ]
        center_c = {1, 2, 3, 4, 5}
        right_c  = set(range(6, 24))
        num_fmt  = {
            6: "#,##0", 7: "#,##0", 8: "#,##0", 9: "#,##0", 10: "#,##0",
            11: "#,##0", 12: "#,##0", 13: "#,##0", 14: "#,##0",
            15: "#,##0", 16: "#,##0", 17: "#,##0",
            18: "0.0", 19: "0.0", 20: "0.0", 21: "0.0", 22: "0.0", 23: "#,##0"
        }
        hl_c = {18, 19, 20, 21, 22}

        _write_data_row(ws2, r2, row_data, s, r2 % 2 == 0, center_c, right_c, num_fmt, hl_c)
        r2 += 1

    # 합계 행
    _write_total_row(ws2, r2, f"전체 합계 ({len(user_summaries)}명)", 5,
                     [(6, "#,##0", False), (7, "#,##0", False), (8, "#,##0", False),
                      (9, "#,##0", False), (10, "#,##0", False),
                      (11, "#,##0", False), (12, "#,##0", False), (13, "#,##0", False), (14, "#,##0", False),
                      (15, "#,##0", False), (16, "#,##0", False), (17, "#,##0", False),
                      (18, "0.0", False), (19, "0.0", False), (20, "0.0", False), (21, "0.0", False),
                      (22, "0.0", True), (23, "#,##0", False)], s)

    # ─────────────────────────────────────────────
    # 시트 3: 특근신청_전체원장
    # ─────────────────────────────────────────────
    ws3 = wb.create_sheet(title="특근신청_전체원장")
    _write_title(ws3,
                 "특근(초과근무) 원 신청서 등록 원장 내역 (4단계 진행상태 포함)",
                 f"출력일시: {now_str}  |  총 {len(records)}건  |  단계: 신청/승인/확정/검토완료",
                 22, s)

    headers3 = [
        "순번", "사번", "성명", "소속팀", "특근분류", "보너스 부여",
        "진행단계", "시작일", "종료일", "일수", "대체휴무 사용일", "대체휴가 사용일수",
        "사전차감", "출장기간(시작)", "출장기간(종료)",
        "프로젝트 번호", "근무 장소", "특근 사유", "승인(승인자)", "확정(확정자)", "검토(검토자)", "신청일시"
    ]
    fills3 = [None] * len(headers3)
    fills3[6] = s["fill_slate"]
    _write_header_row(ws3, 3, headers3, fills3, s)
    ws3.freeze_panes = "A4"

    r3 = 4
    for idx, r in enumerate(records, 1):
        try:
            d1 = datetime.strptime(r["start_date"], "%Y-%m-%d")
            d2 = datetime.strptime(r.get("end_date") or r["start_date"], "%Y-%m-%d")
            days_count = max(1, (d2 - d1).days + 1)
        except Exception:
            days_count = 1

        is_conf = bool(r.get("is_confirmed", 0))
        is_fin = bool(r.get("is_finalized", 0))
        is_rev = bool(r.get("is_reviewed", 0))

        if is_rev:
            stage_text = "검토완료"
        elif is_fin:
            stage_text = "확정"
        elif is_conf:
            stage_text = "승인"
        else:
            stage_text = "신청"

        sub_date = r.get("sub_holiday_date", "") or "-"
        sub_used = float(r.get("sub_holiday_used", 0.0) or 0.0)
        bonus_val = "부여(O)" if r.get("bonus_granted", 0) == 1 else "미부여(-)"
        pre_deduct_val = "O" if int(r.get("is_pre_deduct", 0) or 0) else "-"
        trip_s = r.get("trip_start_date", "") or "-"
        trip_e = r.get("trip_end_date", "") or "-"

        fin_info = f"{r.get('finalized_by', '')}" if is_fin and r.get("finalized_by") else ("완료" if is_fin else "-")
        rev_info = f"{r.get('reviewed_by', '')}" if is_rev and r.get("reviewed_by") else ("완료" if is_rev else "-")
        conf_info = r.get("confirmed_by", "") or ("승인완료" if is_conf else "승인대기")

        row_data = [
            idx,
            r.get("emp_id", ""),
            r.get("user_name", ""),
            r.get("team", ""),
            r.get("category", ""),
            bonus_val,
            stage_text,
            r.get("start_date", ""),
            r.get("end_date", "") or r.get("start_date", ""),
            int(days_count),
            sub_date,
            sub_used,
            pre_deduct_val,
            trip_s,
            trip_e,
            r.get("project_no", "") or "-",
            r.get("location", "") or "-",
            r.get("reason", "") or "-",
            conf_info,
            fin_info,
            rev_info,
            r.get("created_at", "")
        ]
        center_c = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 13, 14, 15, 19, 20, 21, 22}
        right_c  = {10, 12}
        num_fmt  = {10: "#,##0", 12: "0.0"}

        _write_data_row(ws3, r3, row_data, s, r3 % 2 == 0, center_c, right_c, num_fmt)

        stage_cell = ws3.cell(row=r3, column=7)
        _apply_stage_cell_style(stage_cell, stage_text, s)

        r3 += 1

    _write_total_row(ws3, r3, f"합계 ({len(records)}건)", 9,
                     [(10, "#,##0", False), (12, "0.0", False)], s)
    ws3.cell(row=r3, column=11, value="-").alignment = s["align_center"]
    ws3.cell(row=r3, column=11).fill = s["fill_total"]
    ws3.cell(row=r3, column=11).border = s["border_total"]
    for c in range(13, 23):
        ws3.cell(row=r3, column=c, value="").fill = s["fill_total"]
        ws3.cell(row=r3, column=c).border = s["border_total"]

    # ─────────────────────────────────────────────
    # 시트 4: 부서별_요약 (4단계 일수 포함)
    # ─────────────────────────────────────────────
    ws4 = wb.create_sheet(title="부서별_요약")
    _write_title(ws4,
                 "부서(소속팀)별 특근 휴일 현황 및 4단계 구분 총괄표",
                 f"취합 기간: {period_str}  |  출력 일시: {now_str}  |  단계: 신청 / 승인 / 확정 / 검토완료",
                 19, s)

    headers4 = [
        "순번", "소속팀", "소속 인원수", "대체근무 (일)", "법정휴일 (일)",
        "일반휴일 (일)", "대체휴무 (일)", "총 특근일수\n(대체+법정+일반)",
        "신청 (일)", "승인 (일)", "확정 (일)", "검토완료 (일)",
        "보너스 부여 (건)",
        "신청: 실특근+보너스 [일]", "승인: 실특근+보너스 [일]",
        "확정: 실특근+보너스 [일]", "검토완료: 실특근+보너스 [일]",
        "★ 팀 최종 실특근일+보너스 [일]", "총 신청건수"
    ]
    fills4 = [
        None, None, None, s["fill_slate"], s["fill_slate"],
        s["fill_navy"], s["fill_navy"], s["fill_navy"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_gold"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_navy"], None
    ]
    _write_header_row(ws4, 3, headers4, fills4, s)
    ws4.freeze_panes = "A4"

    team_summaries = aggregate_team_holidays(user_summaries)
    r4 = 4
    for idx, tm in enumerate(team_summaries, 1):
        row_data = [
            idx,
            tm["team"],
            int(tm["member_count"]),
            int(tm["sub_work_days"]),
            int(tm["legal_holiday_days"]),
            int(tm["normal_holiday_days"]),
            int(tm.get("sub_holiday_days", 0)),
            int(tm["total_days"]),
            int(tm.get("applied_days", 0)),
            int(tm.get("approved_days", 0)),
            int(tm.get("finalized_days", 0)),
            int(tm.get("reviewed_days", 0)),
            int(tm.get("bonus_count", 0)),
            float(tm.get("applied_final_with_bonus", 0.0)),
            float(tm.get("approved_final_with_bonus", 0.0)),
            float(tm.get("finalized_final_with_bonus", 0.0)),
            float(tm.get("reviewed_final_with_bonus", 0.0)),
            float(tm.get("actual_overtime_with_bonus", 0.0)),
            int(tm["records_count"])
        ]
        center_c = {1, 2}
        right_c  = set(range(3, 20))
        num_fmt  = {
            3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "#,##0",
            9: "#,##0", 10: "#,##0", 11: "#,##0", 12: "#,##0", 13: "#,##0",
            14: "0.0", 15: "0.0", 16: "0.0", 17: "0.0", 18: "0.0", 19: "#,##0"
        }
        hl_c = {14, 15, 16, 17, 18}

        _write_data_row(ws4, r4, row_data, s, r4 % 2 == 0, center_c, right_c, num_fmt, hl_c)
        r4 += 1

    _write_total_row(ws4, r4, f"합계 ({len(team_summaries)}개 팀)", 2,
                     [(3, "#,##0", False), (4, "#,##0", False), (5, "#,##0", False),
                      (6, "#,##0", False), (7, "#,##0", False), (8, "#,##0", False),
                      (9, "#,##0", False), (10, "#,##0", False), (11, "#,##0", False), (12, "#,##0", False),
                      (13, "#,##0", False),
                      (14, "0.0", False), (15, "0.0", False), (16, "0.0", False), (17, "0.0", False),
                      (18, "0.0", True), (19, "#,##0", False)], s)

    # 전 시트 열 너비 자동 조정
    for ws in [ws1, ws2, ws3, ws4]:
        _auto_col_width(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ===========================================================================
#  메인 엑셀 생성 함수 2: 실특근 정산표 전용 (서버 API 사용)
# ===========================================================================

def generate_settlement_excel(
    summary_data: dict,
    start_date: str = "",
    end_date: str = "",
    team_filter: str = ""
) -> io.BytesIO:
    """
    실특근 정산 요약 데이터를 openpyxl로 4단계 라이프사이클(신청/승인/확정/검토) 포함 워크북 생성
    - 시트1: [개인별_실특근_정산표]  인원별 분류별 일수 + 4단계(신청/승인/확정/검토) 일수 + 실특근 집계
    - 시트2: [부서별_정산_요약표]    팀별 현황 + 4단계 일수
    """
    s = _make_styles()
    wb = Workbook()
    now_str = datetime.now().strftime("%Y년 %m월 %d일 %H:%M")

    period_str = ""
    if start_date and end_date:
        period_str = f"조회기간: {start_date} ~ {end_date}"
    elif start_date:
        period_str = f"조회기간: {start_date} 이후"
    elif end_date:
        period_str = f"조회기간: {end_date} 이전"
    else:
        period_str = "조회기간: 전체"

    if team_filter:
        period_str += f"  |  소속팀: {team_filter}"

    users = summary_data.get("user_summary", [])
    teams = summary_data.get("team_summary", [])

    # ─────────────────────────────────────────────
    # 시트 1: 개인별_실특근_정산표
    # ─────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "개인별_실특근_정산표"

    _write_title(ws1,
                 "★ 개인별 최종 실특근일 정산표 (4단계 진행구분)",
                 f"{period_str}  |  4단계: 신청 → 승인 → 확정 → 검토완료  |  출력일시: {now_str}",
                 22, s)

    headers1 = [
        "순번", "사원번호", "성명", "소속팀",
        "대체근무 (일)", "법정휴일 (일)", "일반휴일 (일)", "대체휴무 (일)",
        "총 특근일수 (일)",
        "신청단계 (일)", "승인단계 (일)", "확정단계 (일)", "검토완료 (일)",
        "사전차감 (회)", "사전차감 잔여 (회)", "보너스 부여 (건)",
        "신청: 실특근+보너스 [일]", "승인: 실특근+보너스 [일]",
        "확정: 실특근+보너스 [일]", "검토완료: 실특근+보너스 [일]",
        "★ 최종 실특근일+보너스 [일]", "신청건수"
    ]
    fills1 = [
        None, None, None, None,
        s["fill_slate"], s["fill_slate"], s["fill_navy"], s["fill_navy"],
        s["fill_navy"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_teal"], s["fill_teal"], s["fill_gold"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_navy"], None
    ]
    _write_header_row(ws1, 3, headers1, fills1, s)
    ws1.freeze_panes = "A4"

    r1 = 4
    for idx, u in enumerate(users, 1):
        sub_d = int(u.get("excluded_sub_days", 0) or 0)
        leg_d = int(u.get("excluded_legal_days", 0) or 0)
        ot_d = int(u.get("overtime_days", 0) or 0)
        calculated_total = sub_d + leg_d + ot_d
        actual_ot = float(u.get("actual_overtime_days", 0.0) or 0.0)
        b_cnt = int(u.get("bonus_count", 0) or 0)
        actual_with_b = float(u.get("actual_overtime_with_bonus", actual_ot + b_cnt))

        row_data = [
            idx,
            u.get("emp_id", ""),
            u.get("name", ""),
            u.get("team", ""),
            sub_d,
            leg_d,
            ot_d,
            int(u.get("sub_holiday_days", 0) or u.get("sub_holiday_used", 0) or 0),
            calculated_total,
            int(u.get("applied_days", 0) or 0),
            int(u.get("approved_days", 0) or 0),
            int(u.get("finalized_days", 0) or 0),
            int(u.get("reviewed_days", 0) or 0),
            int(u.get("pre_deduct_count", 0) or u.get("pre_deduct_days", 0) or 0),
            int(u.get("pre_deduct_remaining", 0) or 0),
            b_cnt,
            float(u.get("applied_final_with_bonus", 0.0) or 0.0),
            float(u.get("approved_final_with_bonus", 0.0) or 0.0),
            float(u.get("finalized_final_with_bonus", 0.0) or 0.0),
            float(u.get("reviewed_final_with_bonus", 0.0) or 0.0),
            actual_with_b,
            int(u.get("records_count", 1) or 1)
        ]
        center_c = {1, 2, 3, 4}
        right_c  = set(range(5, 23))
        num_fmt  = {
            5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "#,##0", 9: "#,##0",
            10: "#,##0", 11: "#,##0", 12: "#,##0", 13: "#,##0",
            14: "#,##0", 15: "#,##0", 16: "#,##0",
            17: "0.0", 18: "0.0", 19: "0.0", 20: "0.0", 21: "0.0", 22: "#,##0"
        }
        hl_c = {17, 18, 19, 20, 21}

        _write_data_row(ws1, r1, row_data, s, r1 % 2 == 0, center_c, right_c, num_fmt, hl_c)
        r1 += 1

    # 합계 행
    _write_total_row(ws1, r1, f"전체 합계 ({len(users)}명)", 4,
                     [(5, "#,##0", False), (6, "#,##0", False), (7, "#,##0", False),
                      (8, "#,##0", False), (9, "#,##0", False),
                      (10, "#,##0", False), (11, "#,##0", False), (12, "#,##0", False), (13, "#,##0", False),
                      (14, "#,##0", False), (15, "#,##0", False), (16, "#,##0", False),
                      (17, "0.0", False), (18, "0.0", False), (19, "0.0", False), (20, "0.0", False),
                      (21, "0.0", True), (22, "#,##0", False)], s)

    # ─────────────────────────────────────────────
    # 시트 2: 부서별_정산_요약표
    # ─────────────────────────────────────────────
    ws2 = wb.create_sheet(title="부서별_정산_요약표")
    _write_title(ws2,
                 "부서(소속팀)별 특근 정산 요약표 (4단계 구분)",
                 f"{period_str}  |  팀별 인원수 및 4단계 일수 합산  |  출력일시: {now_str}",
                 17, s)

    headers2 = [
        "순번", "소속팀", "소속 인원수", "총 신청일수",
        "제외일수 (대체+법정)", "인정 특근일 (일반휴일)",
        "대체휴가 사용일수",
        "신청 (일)", "승인 (일)", "확정 (일)", "검토완료 (일)",
        "보너스 부여 (건)",
        "신청: 실특근+보너스 [일]", "승인: 실특근+보너스 [일]",
        "확정: 실특근+보너스 [일]", "검토완료: 실특근+보너스 [일]",
        "★ 팀 최종 실특근일+보너스 [일]"
    ]
    fills2 = [
        None, None, None, s["fill_slate"],
        s["fill_slate"], s["fill_navy"], s["fill_navy"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_gold"],
        s["fill_header_apply"], s["fill_slate"], s["fill_teal"], s["fill_purple"],
        s["fill_navy"]
    ]
    _write_header_row(ws2, 3, headers2, fills2, s)
    ws2.freeze_panes = "A4"

    r2 = 4
    for idx, t in enumerate(teams, 1):
        actual_team_ot = float(t.get("actual_overtime_days", 0.0) or 0.0)
        team_bonus_cnt = int(t.get("bonus_count", 0) or 0)
        actual_team_with_b = float(t.get("actual_overtime_with_bonus", actual_team_ot + team_bonus_cnt))
        row_data = [
            idx,
            t.get("team", ""),
            int(t.get("member_count", 0) or 0),
            int(t.get("total_days", 0) or 0),
            int(t.get("excluded_days", 0) or 0),
            int(t.get("overtime_days", 0) or 0),
            float(t.get("sub_holiday_used", 0.0) or 0.0),
            int(t.get("applied_days", 0) or 0),
            int(t.get("approved_days", 0) or 0),
            int(t.get("finalized_days", 0) or 0),
            int(t.get("reviewed_days", 0) or 0),
            team_bonus_cnt,
            float(t.get("applied_final_with_bonus", 0.0) or 0.0),
            float(t.get("approved_final_with_bonus", 0.0) or 0.0),
            float(t.get("finalized_final_with_bonus", 0.0) or 0.0),
            float(t.get("reviewed_final_with_bonus", 0.0) or 0.0),
            actual_team_with_b
        ]
        center_c = {1, 2}
        right_c  = set(range(3, 18))
        num_fmt  = {
            3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "0.0",
            8: "#,##0", 9: "#,##0", 10: "#,##0", 11: "#,##0",
            12: "#,##0",
            13: "0.0", 14: "0.0", 15: "0.0", 16: "0.0", 17: "0.0"
        }
        hl_c = {13, 14, 15, 16, 17}

        _write_data_row(ws2, r2, row_data, s, r2 % 2 == 0, center_c, right_c, num_fmt, hl_c)
        r2 += 1

    _write_total_row(ws2, r2, f"합계 ({len(teams)}개 팀)", 2,
                     [(3, "#,##0", False), (4, "#,##0", False), (5, "#,##0", False),
                      (6, "#,##0", False), (7, "0.0", False),
                      (8, "#,##0", False), (9, "#,##0", False), (10, "#,##0", False), (11, "#,##0", False),
                      (12, "#,##0", False),
                      (13, "0.0", False), (14, "0.0", False), (15, "0.0", False), (16, "0.0", False),
                      (17, "0.0", True)], s)

    # 전 시트 열 너비 자동 조정
    for ws in [ws1, ws2]:
        _auto_col_width(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def import_overtimes_from_excel(file_bytes: bytes, target_teams: list = None, admin_emp_id: str = None) -> dict:
    """
    엑셀 파일(특근현황_일자별개인별정산의 1번째 시트 또는 특근목록)에서 데이터를 읽어와
    선택된 부서(target_teams)의 특근 정보를 DB(overtimes)에 갱신/추가합니다.
    
    - 선택된 부서만 처리하며 선택되지 않은 부서의 정보는 100% 보존합니다.
    - 동일 사원/동일 날짜/동일 구분의 기록은 중복 생성을 방지(기존 데이터 갱신)합니다.
    - 실패/오류 항목은 사유별 상세 로그를 반환합니다.
    """
    import openpyxl
    import re
    from database import get_db_connection, log_audit

    if isinstance(file_bytes, (bytes, bytearray)):
        bio = io.BytesIO(file_bytes)
    else:
        bio = file_bytes

    try:
        wb = openpyxl.load_workbook(bio, data_only=True)
    except Exception as e:
        return {
            "success": False,
            "message": f"엑셀 파일을 열 수 없습니다: {str(e)}",
            "total_rows": 0, "processed_count": 0, "created_count": 0, "updated_count": 0,
            "skipped_count": 0, "ignored_teams_count": 0, "errors": []
        }

    ws = wb.worksheets[0]

    # 1. 헤더 행 찾기
    header_row_idx = None
    headers_map = {}

    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        row_str = " ".join(row_vals)
        if any(kw in row_str for kw in ["신청자", "사번", "성명", "근무기간", "특근일자", "분류", "소속팀"]):
            header_row_idx = r
            for c_idx, val in enumerate(row_vals, 1):
                clean_v = val.replace(" ", "").replace("\n", "")
                headers_map[clean_v] = c_idx
            break

    if not header_row_idx:
        return {
            "success": False,
            "message": "엑셀 파일에서 올바른 헤더(사번/신청자, 소속팀, 근무기간/특근일자 등)를 찾을 수 없습니다.",
            "total_rows": 0, "processed_count": 0, "created_count": 0, "updated_count": 0,
            "skipped_count": 0, "ignored_teams_count": 0, "errors": []
        }

    # 헤더 인덱스 매핑 찾기
    def find_col(possible_names):
        for k, col in headers_map.items():
            for name in possible_names:
                if name in k:
                    return col
        return None

    emp_col = find_col(["신청자", "사번", "사원번호"])
    name_col = find_col(["성명", "이름"])
    team_col = find_col(["소속팀", "소속", "부서"])
    cat_col = find_col(["분류", "특근구분", "구분"])
    period_col = find_col(["근무기간", "특근기간", "특근일자", "일자"])
    start_col = find_col(["시작일", "시작일자"])
    end_col = find_col(["종료일", "종료일자"])
    proj_col = find_col(["프로젝트"])
    loc_col = find_col(["장소", "근무장소"])
    reason_col = find_col(["사유", "특근사유"])
    bonus_col = find_col(["보너스"])
    pre_deduct_col = find_col(["사전차감", "대체휴무"])

    # 필터 타겟 부서 정제
    valid_target_teams = set()
    if target_teams:
        if isinstance(target_teams, str):
            valid_target_teams = set(t.strip() for t in target_teams.split(",") if t.strip())
        elif isinstance(target_teams, (list, set, tuple)):
            valid_target_teams = set(t.strip() for t in target_teams if t and str(t).strip())

    total_rows = 0
    processed_count = 0
    created_count = 0
    updated_count = 0
    skipped_count = 0
    ignored_teams_count = 0
    errors = []

    conn = get_db_connection()
    cursor = conn.cursor()

    # 사원 DB 정보 캐시 (이름/팀 보완용)
    cursor.execute("SELECT emp_id, name, team FROM users")
    db_users = {row["emp_id"]: dict(row) for row in cursor.fetchall()}

    for r in range(header_row_idx + 1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        str_vals = [str(v or "").strip() for v in row_vals]
        if not any(str_vals) or any(str_vals[0].startswith(kw) for kw in ["합계", "총계", "전체"]):
            continue

        total_rows += 1

        # 1. 사번 / 성명 추출
        raw_emp = str_vals[emp_col - 1] if emp_col and emp_col <= len(str_vals) else ""
        raw_name = str_vals[name_col - 1] if name_col and name_col <= len(str_vals) else ""

        emp_id = ""
        name = ""

        if "(" in raw_emp and ")" in raw_emp:
            m = re.search(r"^(.*?)\((.*?)\)", raw_emp)
            if m:
                name = m.group(1).strip()
                emp_id = m.group(2).strip()
        else:
            emp_id = raw_emp.strip()
            name = raw_name.strip()

        if emp_id in db_users:
            if not name:
                name = db_users[emp_id]["name"]
        elif not emp_id and name:
            for uid, uinfo in db_users.items():
                if uinfo["name"] == name:
                    emp_id = uid
                    break

        if not emp_id:
            errors.append({
                "row": r,
                "name": name or "-",
                "emp_id": "-",
                "reason": "사원번호(사번)를 식별할 수 없음"
            })
            skipped_count += 1
            continue

        # 2. 소속팀 추출 및 부서 필터링
        team = str_vals[team_col - 1] if team_col and team_col <= len(str_vals) else ""
        if not team and emp_id in db_users:
            team = db_users[emp_id]["team"]

        if valid_target_teams:
            if team not in valid_target_teams:
                ignored_teams_count += 1
                skipped_count += 1
                continue

        # 3. 날짜 추출
        start_date = ""
        end_date = ""

        if period_col and period_col <= len(str_vals):
            p_val = str_vals[period_col - 1]
            if "~" in p_val:
                parts = p_val.split("~")
                start_date = parts[0].strip()
                end_date = parts[1].strip()
            elif p_val:
                start_date = p_val.strip()
                end_date = p_val.strip()

        if not start_date and start_col and start_col <= len(str_vals):
            raw_s = ws.cell(r, start_col).value
            if isinstance(raw_s, datetime):
                start_date = raw_s.strftime("%Y-%m-%d")
            elif raw_s:
                start_date = str(raw_s).strip().split("T")[0].split(" ")[0]

        if not end_date and end_col and end_col <= len(str_vals):
            raw_e = ws.cell(r, end_col).value
            if isinstance(raw_e, datetime):
                end_date = raw_e.strftime("%Y-%m-%d")
            elif raw_e:
                end_date = str(raw_e).strip().split("T")[0].split(" ")[0]
        elif not end_date:
            end_date = start_date

        def norm_date(d_str):
            if not d_str: return ""
            d_str = d_str.replace(".", "-").replace("/", "-")
            m = re.search(r"(\d{4})[-_](\d{1,2})[-_](\d{1,2})", d_str)
            if m:
                return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
            return d_str

        start_date = norm_date(start_date)
        end_date = norm_date(end_date)

        if not start_date or not end_date:
            errors.append({
                "row": r,
                "name": name,
                "emp_id": emp_id,
                "reason": "근무기간 / 특근일자 날짜 형식이 올바르지 않음"
            })
            skipped_count += 1
            continue

        # 4. 기타 속성 추출
        category = str_vals[cat_col - 1] if cat_col and cat_col <= len(str_vals) else "일반휴일"
        if not category: category = "일반휴일"

        project = str_vals[proj_col - 1] if proj_col and proj_col <= len(str_vals) else ""
        location = str_vals[loc_col - 1] if loc_col and loc_col <= len(str_vals) else ""
        reason = str_vals[reason_col - 1] if reason_col and reason_col <= len(str_vals) else ""

        raw_b = str_vals[bonus_col - 1] if bonus_col and bonus_col <= len(str_vals) else "0"
        bonus_point = 1 if raw_b in ["1", "예", "Y", "True", "O", "1건"] else 0

        raw_p = str_vals[pre_deduct_col - 1] if pre_deduct_col and pre_deduct_col <= len(str_vals) else "0"
        pre_deduct_point = 1 if raw_p in ["1", "예", "Y", "True", "O", "1건"] else 0

        # 5. DB 중복 검사 및 갱신 / 신규 등록
        cursor.execute("""
            SELECT id FROM overtimes
            WHERE emp_id = ? AND start_date = ? AND end_date = ? AND category = ?
        """, (emp_id, start_date, end_date, category))
        exist_row = cursor.fetchone()

        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if exist_row:
            ot_id = exist_row["id"]
            cursor.execute("""
                UPDATE overtimes
                SET user_name = ?, team = ?, project_no = ?, location = ?, reason = ?,
                    bonus_granted = ?, is_pre_deduct = ?, updated_at = ?
                WHERE id = ?
            """, (name, team, project, location, reason, bonus_point, pre_deduct_point, now_str, ot_id))
            updated_count += 1
            log_audit(ot_id, "EXCEL_IMPORT_UPDATE", admin_emp_id or "ADMIN", "관리자", None, {"emp_id": emp_id, "start_date": start_date})
        else:
            cursor.execute("""
                INSERT INTO overtimes (
                    emp_id, user_name, team, category, start_date, end_date,
                    project_no, location, reason, bonus_granted, is_pre_deduct,
                    is_confirmed, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
            """, (emp_id, name, team, category, start_date, end_date, project, location, reason, bonus_point, pre_deduct_point, now_str, now_str))
            ot_id = cursor.lastrowid
            created_count += 1
            log_audit(ot_id, "EXCEL_IMPORT_CREATE", admin_emp_id or "ADMIN", "관리자", None, {"emp_id": emp_id, "start_date": start_date})

        processed_count += 1

    conn.commit()
    conn.close()

    return {
        "success": True,
        "total_rows": total_rows,
        "processed_count": processed_count,
        "created_count": created_count,
        "updated_count": updated_count,
        "skipped_count": skipped_count,
        "ignored_teams_count": ignored_teams_count,
        "errors": errors
    }
