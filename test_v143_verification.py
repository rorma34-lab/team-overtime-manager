import os
import sys
import json
import sqlite3
from pathlib import Path
from datetime import datetime
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from main import app, update_overtime, batch_delete_overtimes
from schemas import OvertimeUpdateRequest
from database import get_db_connection, init_db
from exporter import aggregate_user_holidays, generate_overtime_excel

def test_v143_all():
    print("==================================================")
    print("🚀 TEAM OVERTIME MANAGER v1.43 VERIFICATION TEST")
    print("==================================================")

    # 1. Version check
    assert app.version == "v1.43", f"App version mismatch: expected v1.43, got {app.version}"
    print("✅ 1. FastAPI App Version is v1.43")

    # 2. Setup Test Data
    init_db()
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM overtimes WHERE emp_id = 'test_v143'")
        cursor.execute("DELETE FROM users WHERE emp_id = 'test_v143'")
        cursor.execute("""
            INSERT INTO users (emp_id, name, team, position, is_admin, is_super, created_at)
            VALUES ('test_v143', '검증맨', '품질혁신팀', '책임', 0, 0, '2026-09-01 09:00:00')
        """)

        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Record 1: 대체근무 1일, 보너스 1
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  sub_holiday_date, sub_holiday_used, is_pre_deduct, trip_start_date, trip_end_date, is_confirmed, bonus_granted, created_at, updated_at)
            VALUES ('test_v143', '검증맨', '품질혁신팀', '대체근무', '2026-09-03', '2026-09-03', 'PRJ-143', '연구소', '대체근무 수행',
                    '', 0.0, 0, '', '', 1, 1, ?, ?)
        """, (now_str, now_str))
        ot_id_1 = cursor.lastrowid

        # Record 2: 법정휴일 1일, 보너스 0
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  sub_holiday_date, sub_holiday_used, is_pre_deduct, trip_start_date, trip_end_date, is_confirmed, bonus_granted, created_at, updated_at)
            VALUES ('test_v143', '검증맨', '품질혁신팀', '법정휴일', '2026-09-04', '2026-09-04', 'PRJ-143', '연구소', '법정휴일 근무',
                    '', 0.0, 0, '', '', 1, 0, ?, ?)
        """, (now_str, now_str))
        ot_id_2 = cursor.lastrowid

        # Record 3: 일반휴일 2일, 보너스 1, 사전차감 1
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  sub_holiday_date, sub_holiday_used, is_pre_deduct, trip_start_date, trip_end_date, is_confirmed, bonus_granted, created_at, updated_at)
            VALUES ('test_v143', '검증맨', '품질혁신팀', '일반휴일', '2026-09-05', '2026-09-06', 'PRJ-143', '본사', '주말 일반휴일',
                    '', 0.0, 1, '', '', 1, 1, ?, ?)
        """, (now_str, now_str))
        ot_id_3 = cursor.lastrowid

        # Record 4: 대체휴무 1일, 보너스 0
        cursor.execute("""
            INSERT INTO overtimes (emp_id, user_name, team, category, start_date, end_date, project_no, location, reason,
                                  sub_holiday_date, sub_holiday_used, is_pre_deduct, trip_start_date, trip_end_date, is_confirmed, bonus_granted, created_at, updated_at)
            VALUES ('test_v143', '검증맨', '품질혁신팀', '대체휴무', '2026-09-10', '2026-09-10', 'PRJ-143', '휴무', '대체휴무 사용',
                    '', 0.0, 0, '2026-09-01', '2026-09-15', 1, 0, ?, ?)
        """, (now_str, now_str))
        ot_id_4 = cursor.lastrowid

        conn.commit()

    print(f"✅ 2. Test data inserted: ids = [{ot_id_1}, {ot_id_2}, {ot_id_3}, {ot_id_4}]")

    # 3. Exporter aggregation logic test
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM overtimes WHERE emp_id = 'test_v143'")
        records = [dict(r) for r in cursor.fetchall()]

    summaries = aggregate_user_holidays(records, {"test_v143": "책임"})
    u = next(s for s in summaries if s["emp_id"] == "test_v143")

    # 수식 1) 총 특근일수 = 대체근무(1) + 법정휴일(1) + 일반휴일(2) = 4 (대체휴무 1일은 제외되어야 함!)
    assert u["sub_work_days"] == 1, f"Expected sub_work_days=1, got {u['sub_work_days']}"
    assert u["legal_holiday_days"] == 1, f"Expected legal_holiday_days=1, got {u['legal_holiday_days']}"
    assert u["normal_holiday_days"] == 2, f"Expected normal_holiday_days=2, got {u['normal_holiday_days']}"
    assert u["sub_holiday_days"] == 1, f"Expected sub_holiday_days=1, got {u['sub_holiday_days']}"
    assert u["total_days"] == 4, f"Expected total_days=4 (1+1+2), got {u['total_days']}"

    # 수식 2) 보너스 개수 = 2건
    assert u["bonus_count"] == 2, f"Expected bonus_count=2, got {u['bonus_count']}"

    print(f"✅ 3. Aggregation verified: total_days={u['total_days']} (대체+법정+일반), bonus_count={u['bonus_count']}")

    # 4. Excel Generation & Header Parsing
    excel_buf = generate_overtime_excel(records, {"test_v143": "책임"})
    wb = openpyxl.load_workbook(excel_buf)
    assert "개인별_휴일합산_정산표" in wb.sheetnames, "Sheet 2 missing"
    ws2 = wb["개인별_휴일합산_정산표"]

    headers_row3 = [ws2.cell(row=3, column=col).value for col in range(1, 17)]
    # 중복 '대체휴가 사용일수'가 삭제되었는지 확인
    assert not any("대체휴가 사용일수" in str(h) for h in headers_row3 if h), "Duplicate '대체휴가 사용일수' column still present!"
    # '보너스 부여 (건)' 열이 추가되었는지 확인
    assert any("보너스 부여" in str(h) for h in headers_row3 if h), "'보너스 부여 (건)' column not found in headers!"
    # '★ 최종 실특근일+보너스' 열이 추가되었는지 확인
    assert any("최종 실특근일+보너스" in str(h) for h in headers_row3 if h), "'★ 최종 실특근일+보너스' column not found in headers!"
    print("✅ 4. Excel headers verified: 중복 대체휴가 사용일수 제거, 보너스 부여 (건) & ★ 최종 실특근일+보너스 정상 존재")

    # Row 4 check
    row4_vals = [ws2.cell(row=4, column=col).value for col in range(1, 17)]
    print(f"📊 Sheet 2 Row 4 Data: {row4_vals}")
    # col 10 is 총 특근일수: 4
    assert row4_vals[9] == 4, f"Expected total_days cell = 4, got {row4_vals[9]}"
    # col 13 is ★ 최종 실특근일: 1.0 (일반휴일 2일 - 사전차감 1일 - (대휴 1일 - 출장내차감 1일) = 1.0)
    assert row4_vals[12] == 1.0, f"Expected actual_ot = 1.0, got {row4_vals[12]}"
    # col 14 is 보너스 부여 (건): 2
    assert row4_vals[13] == 2, f"Expected bonus cell = 2, got {row4_vals[13]}"
    # col 15 is ★ 최종 실특근일+보너스: 3.0 (1.0 + 2 = 3.0)
    assert row4_vals[14] == 3.0, f"Expected actual_ot_with_bonus = 3.0, got {row4_vals[14]}"
    print("✅ 5. Excel Row 4 data verified: total_days=4, actual_ot=1.0, bonus=2, actual_ot+bonus=3.0")

    # 5. Admin PUT Overtime with bonus_granted update
    update_req = OvertimeUpdateRequest(
        changed_by="ps37082",
        category="대체근무",
        start_date="2026-09-03",
        end_date="2026-09-03",
        project_no="PRJ-143-MOD",
        location="연구소",
        reason="보너스 회수 테스트",
        sub_holiday_used=0.0,
        sub_holiday_date="",
        is_pre_deduct=0,
        bonus_granted=0,  # 1 -> 0으로 수정
        trip_start_date="",
        trip_end_date=""
    )
    res_update = update_overtime(ot_id_1, update_req)
    assert res_update["overtime"]["bonus_granted"] == 0, f"Expected bonus_granted 0, got {res_update['overtime']['bonus_granted']}"
    assert res_update["overtime"]["project_no"] == "PRJ-143-MOD"
    print("✅ 6. Admin update_overtime with bonus_granted toggle verified successfully")

    # 6. Admin batch_delete_overtimes
    res_batch_del = batch_delete_overtimes({
        "ids": [ot_id_1, ot_id_2],
        "admin_emp_id": "ps37082"
    })
    assert res_batch_del["deleted_count"] == 2, f"Expected deleted 2, got {res_batch_del}"
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM overtimes WHERE id IN (?, ?)", (ot_id_1, ot_id_2))
        assert len(cursor.fetchall()) == 0, "Items were not deleted!"
        # Audit log check (overtime_history)
        import time
        time.sleep(0.5)  # wait for background log queue
        cursor.execute("SELECT action FROM overtime_history WHERE overtime_id IN (?, ?)", (ot_id_1, ot_id_2))
        actions = [r["action"] for r in cursor.fetchall()]
        assert "일괄삭제" in actions, f"Missing 일괄삭제 in audit logs: {actions}"
    print("✅ 7. Admin batch_delete_overtimes & Audit log verified successfully")

    # Clean up test user & remaining records
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM overtimes WHERE emp_id = 'test_v143'")
        cursor.execute("DELETE FROM users WHERE emp_id = 'test_v143'")
        conn.commit()

    print("\n🎉 ALL v1.43 BACKEND, EXPORTER, AND BUSINESS LOGIC TESTS PASSED PERFECTLY!")

if __name__ == "__main__":
    test_v143_all()
